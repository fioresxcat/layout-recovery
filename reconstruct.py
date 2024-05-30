import cv2
from copy import deepcopy
from pathlib import Path
from utils import *
import os
import numpy as np
from docx import Document
from docx import shared
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT
from html.parser import HTMLParser
from shapely.geometry import Polygon
import pdb



def sorted_layout_boxes(res, w):
    """
    Sort text boxes in order from top to bottom, left to right
    args:
        res(list):ppstructure results
    return:
        sorted results(list)
    """
    num_boxes = len(res)
    if num_boxes == 1:
        res[0]['layout'] = 'single'
        return res

    sorted_boxes = sorted(res, key=lambda x: (x['bbox'][1], x['bbox'][0]))
    _boxes = list(sorted_boxes)

    new_res = []
    res_left = []
    res_right = []
    i = 0

    while True:
        if i >= num_boxes:
            break
        if i == num_boxes - 1:
            if _boxes[i]['bbox'][1] > _boxes[i - 1]['bbox'][3] and _boxes[i]['bbox'][0] < w / 2 and _boxes[i]['bbox'][2] > w / 2:
                new_res += res_left
                new_res += res_right
                _boxes[i]['layout'] = 'single'
                new_res.append(_boxes[i])

            else:
                if _boxes[i]['bbox'][2] > w / 2:
                    _boxes[i]['layout'] = 'double'
                    res_right.append(_boxes[i])
                    new_res += res_left
                    new_res += res_right
                elif _boxes[i]['bbox'][0] < w / 2:
                    _boxes[i]['layout'] = 'double'
                    res_left.append(_boxes[i])
                    new_res += res_left
                    new_res += res_right
            res_left = []
            res_right = []
            break
        
        elif _boxes[i]['bbox'][0] < w / 4 and _boxes[i]['bbox'][2] < 3 * w / 4:
            _boxes[i]['layout'] = 'double'
            res_left.append(_boxes[i])
            i += 1

        elif _boxes[i]['bbox'][0] > w / 4 and _boxes[i]['bbox'][2] > w / 2:
            _boxes[i]['layout'] = 'double'
            res_right.append(_boxes[i])
            i += 1

        else:
            new_res += res_left
            new_res += res_right
            _boxes[i]['layout'] = 'single'
            new_res.append(_boxes[i])
            res_left = []
            res_right = []
            i += 1
    if res_left:
        new_res += res_left
    if res_right:
        new_res += res_right
    return new_res



def convert_info_docx(img, res, save_folder, img_name):
    doc = Document()
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    doc.styles['Normal'].font.size = shared.Pt(6.5)

    flag = 1

    for i, region in enumerate(res):
        if len(region['res']) == 0 and region['type'].lower() != 'figure':
            continue
        img_idx = region['img_idx']
        if flag == 2 and region['layout'] == 'single':
            section = doc.add_section(WD_SECTION.CONTINUOUS)
            section._sectPr.xpath('./w:cols')[0].set(qn('w:num'), '1')
            flag = 1
        elif flag == 1 and region['layout'] == 'double':
            section = doc.add_section(WD_SECTION.CONTINUOUS)
            section._sectPr.xpath('./w:cols')[0].set(qn('w:num'), '2')
            flag = 2

        if region['type'].lower() == 'figure':
            img_path = os.path.join(save_folder, img_name, f'figure_{i}.png')
            cv2.imwrite(img_path, region['img'])
            paragraph_pic = doc.add_paragraph()
            paragraph_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph_pic.add_run("")
            if flag == 1:
                run.add_picture(img_path, width=shared.Inches(5))
            elif flag == 2:
                run.add_picture(img_path, width=shared.Inches(2))

        elif region['type'].lower() == 'title':
            doc.add_heading(region['res'][0]['text'])

        elif region['type'].lower() == 'table':
            parser = HtmlToDocx()
            parser.table_style = 'TableGrid'
            parser.handle_table(region['res']['html'], doc)

        else:  # list and text
            paragraph = doc.add_paragraph()
            paragraph_format = paragraph.paragraph_format
            for i, line in enumerate(region['res']):
                if i == 0:
                    paragraph_format.first_line_indent = shared.Inches(0.25)
                text_run = paragraph.add_run(line['text'] + ' ')
                text_run.font.size = shared.Pt(10)

    return doc




class Reconstructor:
    def __init__(self, common_cfg, model_cfg) -> None:
        self.common_cfg = common_cfg
        self.model_cfg = model_cfg

    
    def format_result(self, result):
        """
            format result from ocr and text detection and layout detection into paddle format
        """
        src_img = result['orig_img']
        region_infos = []
        for region_idx, region_bb in enumerate(result['layout']['boxes']):
            region = {}

            score = result['layout']['scores'][region_idx]
            class_name = result['layout']['class_names'][region_idx]

            # get region image
            x1, y1, x2, y2 = region_bb
            region_img = src_img[y1:y2, x1:x2]

            # get region texts
            poly2text = {}
            for poly, text in zip(result['text_detection']['coords'], result['ocr']['raw_words']):
                poly2text[poly] = text

            if class_name in ['text', 'title', 'list']:
                polys = []
                region['res'] = []
                for poly, text in poly2text.items():
                    poly_bb = poly2bb(poly)
                    iou, max_overlap_ratio = compute_boxes_iou(region_bb, poly_bb)
                    if max_overlap_ratio > 0.7:
                        polys.append(poly)
                poly_rows = row_polys(polys)
                for row in poly_rows:
                    row_text = ' '.join([poly2text[poly] for poly in row])
                    row_xmin = min([min(poly[::2]) for poly in row])
                    row_ymin = min([min(poly[1::2]) for poly in row])
                    row_xmax = max([max(poly[::2]) for poly in row])
                    row_ymax = max([max(poly[1::2]) for poly in row])
                    row_coords = [[row_xmin, row_ymin], [row_xmax, row_ymin], [row_xmax, row_ymax], [row_xmin, row_ymax]]
                    region['res'].append({
                        'text': row_text,
                        'text_region': row_coords,
                        'confidence': 1.
                    })

            elif class_name == 'table':
                region['res'] = ''
            
            else:
                region['res'] = ''

            region['type'] = class_name
            region['bbox'] = region_bb
            region['score'] = score
            region['img'] = region_img
            region['img_idx'] = 0
            region_infos.append(region)
        
        return region_infos


    def gather_and_sort_boxes(self, result):
        layout_boxes = result['layout']['boxes']
        layout_texts = [[] for _ in layout_boxes]
        words = result['ocr']['raw_words']
        for i, poly in enumerate(result['text_detection']['coords']):
            for box_idx, box in enumerate(layout_boxes):
                if is_poly_in_box(poly, box):
                    layout_texts[box_idx].append(poly)
                    break
        poly2word = dict(zip(result['text_detection']['coords'], words))
        for i, polys in enumerate(layout_texts):
            poly_rows = row_polys(polys)
            for row in poly_rows:
                for poly_idx, poly in enumerate(row):
                    row[poly_idx] = (poly, poly2word[tuple(poly)])
            layout_texts[i] = poly_rows
        
        result['reconstruct'] = {}
        result['reconstruct']['layout_texts'] = layout_texts
        return result
    

    # def predict(self, result):
    #     img = deepcopy(result['orig_img'])
    #     h, w, _ = img.shape
    #     res = self.format_result(result)
    #     res = sorted_layout_boxes(res, w)
    #     save_dir = os.path.join(self.model_cfg['save_folder'], result['img_name'])
    #     os.makedirs(save_dir, exist_ok=True)
    #     doc = convert_info_docx(img, res, save_folder=self.model_cfg['save_folder'], img_name=result['img_name'])

    #     # visualize result
    #     for i, region in enumerate(res):
    #         x1, y1, x2, y2 = region['bbox']
    #         score = region['score']
    #         cv2.rectangle(img, (x1, y1), (x2, y2), (0, 0, 255), 2)
    #         cv2.putText(img, f'{i}', (x1, y1), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 0, 0), 2)
    #         cx, cy = (x1 + x2) // 2, (y1 + y2) // 2
    #         cv2.putText(img, f'{region["type"]} {score:.2f}', (cx, cy), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 0, 0), 2)

    #     cv2.imwrite(os.path.join(save_dir, 'layout_detection.jpg'), img)
    #     doc.save(os.path.join(save_dir, '{}.docx'.format(result['img_name'])))
    #     result['final_doc'] = doc
    #     return result

    def predict(self, result):
        result = self.gather_and_sort_boxes(result)
        #TODO: hieunt
        return result