from docx import Document
from docx import shared
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
import pdb
import cv2
import fitz
import numpy as np
import os

def main():
    doc = Document()
    # get page width and page height of doc
    for section in doc.sections:
        print(section.page_width)  # Inches(8.5)
        print(section.page_height)  # Inches(11)

    section = doc.add_section(WD_SECTION.CONTINUOUS)
    # set w and h for section
    section.page_width = shared.Inches(2)
    section.page_height = shared.Inches(2)
    # add text to the section
    paragraph = doc.add_paragraph()
    paragraph_format = paragraph.paragraph_format
    text_run = paragraph.add_run(' '.join(['hello world ']*100) + ' ')
    text_run.font.size = shared.Pt(10)
    # set margin for the section
    section.left_margin = shared.Inches(4)
    section.right_margin = shared.Inches(2)
    section.top_margin = shared.Inches(4)
    section.bottom_margin = shared.Inches(2.5)
    # save
    doc.save('test.docx')


def visualize_pdf_blocks():
    doc = fitz.open('output.pdf')
    for page_index, page in enumerate(doc):
        # page: Page = page
        # page.get_textbox((0, 0, page.rect.width, page.rect.height))
        # page.get_textpage()
        # page.get_displaylist()

        mat = fitz.Matrix(2, 2)
        pix = page.get_pixmap(matrix=mat)
        shape = (pix.height, pix.width, 3)
        image = np.ndarray(shape, dtype=np.uint8, buffer=pix.samples)  # this is rgb image
        image = cv2.cvtColor(image, cv2.COLOR_RGB2BGR)  # after this it is actually bgr image

        page_blocks = page.get_text("dict", flags=11)["blocks"]
        # pdb.set_trace()
        for b_idx, block in enumerate(page_blocks):
            bb = list(map(int, block['bbox']))
            bb = list(map(lambda x: x*2, bb))
            cv2.rectangle(image, (bb[0], bb[1]), (bb[2], bb[3]), (0, 0, 255), 2)
        cv2.imwrite(f'page_{page_index}.jpg', image)
        print(f'done page {page_index}')


def nothing():
    from pdf2docx import Converter

    pdf_file = 'output.pdf'
    docx_file = 'output-pdf2docx.docx'

    # convert pdf to docx
    cv = Converter(pdf_file)
    cv.convert(docx_file)      # all pages by default
    cv.close()




if __name__ == '__main__':
    # main()
    visualize_pdf_blocks()